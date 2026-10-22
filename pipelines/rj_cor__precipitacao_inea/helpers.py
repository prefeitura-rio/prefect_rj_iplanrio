"""Helpers of rj_cor__precipitacao_inea."""

from collections.abc import Mapping, Sequence
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from pipelines.rj_cor__precipitacao_inea.constants import (
    MAX_WATER_HEIGHT,
    MISSING,
    PLUVIOMETRIC_COLUMNS,
    TIME_FORMAT,
    WATER_COLUMN,
)


def read_sheet(content: bytes) -> list[dict[str, object]]:
    """Return the rows of the first sheet of an Excel file, by the name of the header cells.

    The names lose the spaces around them.
    """
    sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    names = [str(name).strip() if name is not None else "" for name in next(rows, ())]

    return [dict(zip(names, row, strict=False)) for row in rows]


def read_number(value: object) -> float | None:
    """Return a reading as a number, or nothing if it is missing or is not a number."""
    if value is None or value == MISSING:
        return None

    try:
        return float(str(value))
    except ValueError:
        return None


def read_time(row: Mapping[str, object]) -> str | None:
    """Return the time of a row as `YYYY-MM-DD HH:MM:SS`, or nothing if it does not parse."""
    text = f"{str(row.get('Data', '')).strip()} {str(row.get('Hora', '')).strip()}"

    try:
        return datetime.strptime(text, TIME_FORMAT).strftime("%Y-%m-%d %H:%M:%S")  # noqa: DTZ007  # the sheet has no time zone
    except ValueError:
        return None


def build_rows(
    stations: Mapping[str, Sequence[Mapping[str, object]]],
) -> tuple[list[dict[str, str | float | None]], list[dict[str, str | float | None]]]:
    """Return the pluviometric rows and the fluviometric rows of the sheets of the stations.

    A row without a valid time is dropped, and a repeated station and time keeps the
    first. A water height above 10000 is dropped from the fluviometric rows.
    """
    pluviometric: list[dict[str, str | float | None]] = []
    fluviometric: list[dict[str, str | float | None]] = []

    for station, sheet in stations.items():
        seen: set[str] = set()

        for row in sheet:
            measured = read_time(row)

            if measured is None or measured in seen:
                continue

            seen.add(measured)
            key = {"id_estacao": station, "data_medicao": measured}
            pluviometric.append(
                {**key, **{name: read_number(row.get(source)) for source, name in PLUVIOMETRIC_COLUMNS.items()}}
            )
            height = read_number(row.get(WATER_COLUMN))

            if height is None or height <= MAX_WATER_HEIGHT:
                fluviometric.append({**key, "altura_agua": height})

    return pluviometric, fluviometric
