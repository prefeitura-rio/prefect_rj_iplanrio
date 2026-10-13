# -*- coding: utf-8 -*-
"""Monta e publica os arquivos de uma rodada de relatório, numa subpasta nova da pasta
da campanha (nunca substitui uma rodada anterior — ela pode ter uma planilha revisada):

  <campanha>/<data>_<hora>h__disparos_<inicio>_a_<fim>[__catalogo_automatico]/
      relatorio_engajamento      Google Doc (gerado como .docx e convertido no upload)
      conversas_classificadas    Google Sheet, 1 linha por conversa; no relatório de
                                 engajamento, com as colunas de revisão em branco
                                 (revisao com lista suspensa Certo/Errado)
      catalogo_automatico        Google Sheet, só no relatório rápido"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from googleapiclient.discovery import Resource

from pipelines.rj_crm__relatorio_engajamento_hsm.config import (
    ARQUIVO_CATALOGO_AUTOMATICO,
    ARQUIVO_CONVERSAS,
    ARQUIVO_RELATORIO,
    CATEGORIA_SEM_CATALOGO,
    CHATBOT_DATASET,
    CHATBOT_TABLE,
    COL_CAT_CATEGORIA,
    COL_CAT_DESCRICAO,
    COL_CAT_EXEMPLOS,
    COL_CATEGORIAS_CORRETAS,
    COL_CONV_CATEGORIA,
    COL_CONV_CONVERSA,
    COL_CONV_ID_SESSAO,
    COL_CONV_JUIZ_ESPERADA,
    COL_CONV_JUIZ_JUSTIFICATIVA,
    COL_CONV_JUIZ_VEREDITO,
    COL_CONV_JUSTIFICATIVA,
    COL_CONV_RESUMO,
    COL_JUSTIFICATIVA_REVISAO,
    COL_REVISAO,
    DOCX_MIME,
    DRIVE_DOC_MIME,
    MARCADOR_REVISADA,
    OPCOES_REVISAO,
    PROJECT_ID,
    SUFIXO_PASTA_RAPIDO,
)
from pipelines.rj_crm__relatorio_engajamento_hsm.tasks.discovery import SugestaoCategoria
from pipelines.rj_crm__relatorio_engajamento_hsm.tasks.judge import metricas_por_categoria
from pipelines.rj_crm__relatorio_engajamento_hsm.utils.drive import garante_pasta, link_pasta, upload_bytes
from pipelines.rj_crm__relatorio_engajamento_hsm.utils.sheets import aplica_lista_suspensa, cria_sheet

_Z_IC_95 = 1.959963985  # z crítico do intervalo de confiança de 95%
SEP_CELULA = "; "  # várias categorias numa célula (a ingestão separa por ";" e tira espaços)


def _margem_erro_wilson(p: float, n: int, N: int | None = None, z: float = _Z_IC_95) -> float:
    """Meia-largura do intervalo de Wilson (95%) pra uma proporção binomial p estimada
    com n observações, com correção de população finita (N = tamanho da categoria):
    sem ela, uma categoria julgada 100% (n == N, censo completo, erro amostral zero por
    definição) apareceria com margem de erro > 0, o que é errado."""
    if n <= 0:
        return float("nan")
    if N is not None and n >= N:
        return 0.0
    denom = 1 + z**2 / n
    margem = (z * math.sqrt(p * (1 - p) / n + z**2 / (4 * n**2))) / denom
    if N is not None and N > 1:
        margem *= math.sqrt((N - n) / (N - 1))
    return margem


def texto_precisao(corretas: int, avaliadas: int, total: int) -> str:
    """ "87% (n=46)" se todas as conversas foram julgadas; com amostra (ou lote do juiz
    que falhou), "87% ± 6% (n=46 de 120)"; "não avaliada" se nenhuma."""
    if avaliadas <= 0:
        return "não avaliada"
    p = corretas / avaliadas
    if avaliadas >= total:
        return f"{p:.0%} (n={avaliadas})"
    return f"{p:.0%} ± {_margem_erro_wilson(p, avaliadas, total):.0%} (n={avaliadas} de {total})"


def nome_subpasta(gerado_em: datetime, data_inicio: date, data_fim: date, rapido: bool) -> str:
    nome = f"{gerado_em:%Y-%m-%d_%H}h__disparos_{data_inicio.isoformat()}_a_{data_fim.isoformat()}"
    return nome + (SUFIXO_PASTA_RAPIDO if rapido else "")


# ============================================================================
# Dados da rodada
# ============================================================================


@dataclass
class DadosRelatorio:
    campanha: str
    data_inicio: date
    data_fim: date
    gerado_em: datetime
    rapido: bool
    df: pd.DataFrame  # saída de julga_classificacoes
    descricoes: dict[str, str]  # categoria -> descrição (catálogo usado)
    textos_hsm: dict[str, str]
    total_disparos: int
    total_engajados: int
    sql_utilizada: str
    n_exemplos_rag: int = 0  # exemplos revisados disponíveis no banco (0 = zero-shot)
    categorias_sem_decisao: list[str] = field(default_factory=list)
    exemplos_descoberta: dict[str, list[str]] = field(default_factory=dict)  # só relatório rápido
    nome_eixo: str | None = None
    # Sugestões de categorias novas (só relatório de engajamento). None = não rodou
    # (relatório rápido ou sugerir_categorias desligado); [] = rodou e não achou nada.
    sugestoes: list[SugestaoCategoria] | None = None
    sugestoes_erro: str = ""


@dataclass
class Exemplo:
    conversa: str
    resumo: str


@dataclass
class Categoria:
    nome: str
    descricao: str
    qtd: int
    pct: float
    n_avaliadas: int
    n_corretas: int
    exemplos: list[Exemplo]


def monta_categorias(dados: DadosRelatorio, n_exemplos_por_categoria: int) -> list[Categoria]:
    """Métricas por categoria + exemplos (prioriza conversas que o juiz confirmou)."""
    metricas = metricas_por_categoria(dados.df)
    classificadas = dados.df[dados.df["falha_classificacao"] == ""]
    categorias = []
    for linha in metricas.itertuples():
        com_categoria = classificadas[classificadas["categorias"].map(lambda cats, c=linha.categoria: c in cats)]
        ordenadas = com_categoria.assign(_ok=com_categoria["juiz_veredito"] == "CORRETO").sort_values(
            "_ok", ascending=False, kind="stable"
        )
        descricao = dados.descricoes.get(linha.categoria, "")
        if linha.categoria == CATEGORIA_SEM_CATALOGO:
            descricao = "Conversas que não se encaixaram em nenhuma categoria do catálogo — podem indicar um tema novo."
        categorias.append(
            Categoria(
                nome=linha.categoria,
                descricao=descricao,
                qtd=int(linha.qtd),
                pct=float(linha.pct_do_engajamento),
                n_avaliadas=int(linha.n_avaliadas_juiz),
                n_corretas=int(linha.n_corretas),
                exemplos=[
                    Exemplo(r.conversa_completa, r.resumo_gerado)
                    for r in ordenadas.head(n_exemplos_por_categoria).itertuples()
                ],
            )
        )
    return categorias


# ============================================================================
# Planilhas
# ============================================================================


def monta_conversas_df(df: pd.DataFrame, com_revisao: bool) -> pd.DataFrame:
    """conversas_classificadas: 1 linha por conversa (inclusive as com falha de
    classificação — categoria vazia, motivo na justificativa)."""
    linhas = []
    for r in df.itertuples():
        justificativa = (
            r.categoria_justificativa
            if not r.falha_classificacao
            else (f"falha na classificação: {r.falha_classificacao}")
        )
        linha = {
            COL_CONV_ID_SESSAO: r.id_sessao_48h,
            COL_CONV_CATEGORIA: SEP_CELULA.join(r.categorias),
            COL_CONV_JUSTIFICATIVA: justificativa,
            COL_CONV_RESUMO: r.resumo_gerado,
            COL_CONV_CONVERSA: r.conversa_completa,
            COL_CONV_JUIZ_VEREDITO: r.juiz_veredito,
            COL_CONV_JUIZ_ESPERADA: SEP_CELULA.join(r.juiz_categorias_esperadas),
            COL_CONV_JUIZ_JUSTIFICATIVA: r.juiz_justificativa,
        }
        if com_revisao:
            linha |= {COL_REVISAO: "", COL_CATEGORIAS_CORRETAS: "", COL_JUSTIFICATIVA_REVISAO: ""}
        linhas.append(linha)
    return pd.DataFrame(linhas)


def monta_catalogo_automatico_df(dados: DadosRelatorio, categorias: list[Categoria]) -> pd.DataFrame:
    qtd = {c.nome: c.qtd for c in categorias}
    return pd.DataFrame(
        [
            {
                COL_CAT_CATEGORIA: nome,
                COL_CAT_DESCRICAO: descricao,
                COL_CAT_EXEMPLOS: SEP_CELULA.join(dados.exemplos_descoberta.get(nome, [])),
                "Qtd. de conversas": qtd.get(nome, 0),
            }
            for nome, descricao in dados.descricoes.items()
        ]
    )


# ============================================================================
# Documento
# ============================================================================

_COR_ACCENT = RGBColor(0xB5, 0x50, 0x2E)
_COR_AVISO = RGBColor(0xB0, 0x1E, 0x1E)
_COR_FUNDO_CODIGO = "FBE9E2"
_COR_BORDA_CODIGO = "B5502E"
_MAX_CHARS_EXEMPLO = 1200


def _trunca(texto: str, limite: int = _MAX_CHARS_EXEMPLO) -> str:
    texto = texto or ""
    return texto if len(texto) <= limite else texto[:limite].rstrip() + " […]"


def _sombreia_paragrafo(paragraph, cor_fundo: str, cor_borda: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_fundo)
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        borda = OxmlElement(f"w:{lado}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "4")
        borda.set(qn("w:space"), "4")
        borda.set(qn("w:color"), cor_borda)
        pBdr.append(borda)
    pPr.append(pBdr)


def _adiciona_bloco_codigo(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    linhas = texto.split("\n")
    for i, linha in enumerate(linhas):
        run = p.add_run(linha if linha.strip() else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        if i < len(linhas) - 1:
            run.add_break()
    _sombreia_paragrafo(p, _COR_FUNDO_CODIGO, _COR_BORDA_CODIGO)


def observacoes(dados: DadosRelatorio) -> list[str]:
    """Itens da seção "Observações importantes" (numerados no documento)."""
    itens = [
        "Este relatório foi gerado a partir de uma LLM (um modelo de linguagem) e portanto pode conter erros. "
        "Revise os números com atenção e se atente às métricas e exemplos fornecidos para não tirar conclusões "
        "erradas.",
        "Uma conversa pode ter mais de uma categoria, quando o cidadão trata de assuntos distintos. Por isso a soma "
        "das porcentagens por categoria pode passar de 100%.",
        f"Conversas em {CATEGORIA_SEM_CATALOGO} não se encaixaram em nenhuma categoria do catálogo e podem indicar "
        "um tema que ainda não tem categoria.",
    ]
    if dados.sugestoes:
        itens.append(
            f"A seção 'Sugestões de novas categorias', no fim deste relatório, traz {len(dados.sugestoes)} tema(s) "
            "que apareceram nas conversas e não têm categoria no Catálogo. Eles não foram usados nesta classificação."
        )
    if dados.categorias_sem_decisao:
        itens.append(
            f"O Catálogo desta campanha tinha {len(dados.categorias_sem_decisao)} categoria(s) sem decisão "
            f"({', '.join(dados.categorias_sem_decisao)}), que não foram usadas nesta classificação. Conversas desses "
            "temas podem ter sido classificadas em outra categoria ou ficado sem categoria."
        )
    if dados.rapido:
        itens.append(
            f"As planilhas '{ARQUIVO_CONVERSAS}' e '{ARQUIVO_CATALOGO_AUTOMATICO}', na mesma pasta, trazem a "
            "classificação de cada conversa e as categorias criadas automaticamente. Elas não são usadas como "
            "exemplos em classificações futuras."
        )
    else:
        itens.append(
            f"A planilha '{ARQUIVO_CONVERSAS}', na mesma pasta, traz a classificação de cada conversa. Para revisar: "
            f"preencha '{COL_REVISAO}' (Certo/Errado), '{COL_CATEGORIAS_CORRETAS}' (quando Errado, separadas por ';') "
            f"e '{COL_JUSTIFICATIVA_REVISAO}', e renomeie o arquivo para "
            f"'{ARQUIVO_CONVERSAS}_{MARCADOR_REVISADA}_AAAA-MM-DD'. As correções passam a ser usadas como exemplos "
            "nas próximas classificações desta campanha. Uma categoria nova precisa ser incluída antes no Catálogo."
        )
    return itens


def gera_docx(dados: DadosRelatorio, categorias: list[Categoria]) -> Document:
    df = dados.df
    classificadas = df[df["falha_classificacao"] == ""]
    n_classificadas, n_falhas = len(classificadas), len(df) - len(classificadas)
    julgadas = classificadas[classificadas["juiz_veredito"] != ""]

    doc = Document()
    doc.add_heading(f"Relatório de Engajamento — {dados.campanha}", level=0)
    meta = doc.add_paragraph()
    meta.add_run(
        f"Disparos de {dados.data_inicio:%d/%m/%Y} a {dados.data_fim:%d/%m/%Y}"
        + (f"    ·    Eixo: {dados.nome_eixo}" if dados.nome_eixo else "")
    ).italic = True
    doc.add_paragraph(f"Gerado em {dados.gerado_em:%d/%m/%Y %H:%M}")

    if dados.rapido:
        aviso = doc.add_paragraph().add_run(
            "Relatório rápido: as categorias foram criadas automaticamente pela LLM a partir das conversas deste "
            "período e NÃO foram validadas pela equipe responsável. Use para uma leitura exploratória."
        )
        aviso.bold = True
        aviso.font.color.rgb = _COR_AVISO

    doc.add_heading("Resumo", level=1)
    pct = lambda parte, todo: 100 * parte / todo if todo else 0.0  # noqa: E731
    doc.add_paragraph(f"Sessões que receberam o HSM: {dados.total_disparos}")
    doc.add_paragraph(
        f"Sessões com resposta do cidadão: {dados.total_engajados} "
        f"({pct(dados.total_engajados, dados.total_disparos):.1f}% das sessões)"
    )
    doc.add_paragraph(
        f"Conversas classificadas: {n_classificadas} ({pct(n_classificadas, dados.total_engajados):.1f}% das respostas)"
        + (f" — {n_falhas} com falha na classificação" if n_falhas else "")
    )
    doc.add_paragraph(f"Categorias com conversas: {len(categorias)}")
    if dados.rapido:
        modo = "catálogo criado automaticamente (relatório rápido)"
    elif dados.n_exemplos_rag:
        modo = f"Catálogo validado + {dados.n_exemplos_rag} conversa(s) revisada(s) desta campanha como exemplos"
    else:
        modo = "Catálogo validado, sem exemplos (a campanha ainda não tem conversas revisadas)"
    doc.add_paragraph(f"Classificação: {modo}")
    doc.add_paragraph(
        f"Conversas avaliadas pelo juiz: {len(julgadas)} ({pct(len(julgadas), n_classificadas):.1f}% das classificadas)"
    )
    precisao = doc.add_paragraph()
    precisao.add_run("Precisão da classificação (conversas com o conjunto de categorias correto): ")
    n_corretas = int((julgadas["juiz_veredito"] == "CORRETO").sum())
    precisao.add_run(texto_precisao(n_corretas, len(julgadas), n_classificadas)).bold = True

    if dados.textos_hsm:
        doc.add_paragraph().add_run("HSM(s) enviado(s) (texto do template):").bold = True
        for nome, texto in dados.textos_hsm.items():
            doc.add_paragraph(nome, style="List Bullet")
            doc.add_paragraph(style="Intense Quote").add_run(_trunca(texto))

    doc.add_heading("Observações importantes", level=1)
    for i, item in enumerate(observacoes(dados), start=1):
        doc.add_paragraph(f"{i}) {item}")

    doc.add_heading("Query utilizada na análise", level=1)
    doc.add_paragraph(
        f"Consulta em {PROJECT_ID}.{CHATBOT_DATASET}.{CHATBOT_TABLE} usada pra extrair as conversas, com os "
        "parâmetros já substituídos pelos valores reais desta análise."
    )
    _adiciona_bloco_codigo(doc, dados.sql_utilizada)

    doc.add_heading("Categorias", level=1)
    tabela = doc.add_table(rows=1, cols=4)
    tabela.style = "Light Grid Accent 2"
    for cel, texto in zip(
        tabela.rows[0].cells, ["Categoria", "Qtd.", "% das respostas", "Precisão (LLM juiz)"], strict=True
    ):
        cel.paragraphs[0].add_run(texto).bold = True
    for cat in categorias:
        linha = tabela.add_row().cells
        linha[0].text = cat.nome
        linha[1].text = str(cat.qtd)
        linha[2].text = f"{cat.pct:.1f}%"
        linha[3].text = texto_precisao(cat.n_corretas, cat.n_avaliadas, cat.qtd)
    doc.add_paragraph()

    for cat in categorias:
        h = doc.add_heading(cat.nome, level=2)
        h.runs[0].font.color.rgb = _COR_ACCENT
        doc.add_paragraph(cat.descricao or "(sem descrição no catálogo)")
        resumo = doc.add_paragraph()
        resumo.add_run(f"{cat.qtd} conversa(s) · {cat.pct:.1f}% das respostas · precisão: ")
        resumo.add_run(texto_precisao(cat.n_corretas, cat.n_avaliadas, cat.qtd)).bold = True
        if cat.exemplos:
            doc.add_paragraph().add_run("Exemplos:").bold = True
            tabela_ex = doc.add_table(rows=1, cols=2)
            tabela_ex.style = "Light Grid Accent 2"
            for cel, texto in zip(tabela_ex.rows[0].cells, ["Conversa", "Resumo"], strict=True):
                cel.paragraphs[0].add_run(texto).bold = True
            for ex in cat.exemplos:
                linha_ex = tabela_ex.add_row().cells
                linha_ex[0].text = _trunca(ex.conversa)
                linha_ex[1].text = _trunca(ex.resumo) if ex.resumo else "(sem resumo)"
            doc.add_paragraph()
        else:
            doc.add_paragraph().add_run("Sem exemplo disponível.").italic = True

    if dados.sugestoes is not None or dados.sugestoes_erro:
        _secao_sugestoes(doc, dados)
    return doc


def _secao_sugestoes(doc: Document, dados: DadosRelatorio) -> None:
    """Categorias que a descoberta (partindo do Catálogo) encontrou e o Catálogo não
    tem. Só sugestão: não entram no Catálogo nem na classificação desta rodada."""
    doc.add_heading("Sugestões de novas categorias", level=1)
    if dados.sugestoes_erro:
        doc.add_paragraph().add_run(
            f"Não foi possível gerar sugestões nesta rodada ({dados.sugestoes_erro})."
        ).italic = True
        return
    if not dados.sugestoes:
        doc.add_paragraph(
            "Nenhuma categoria nova sugerida: os temas das conversas deste período já estão cobertos pelas "
            "categorias do Catálogo."
        )
        return
    doc.add_paragraph(
        "A LLM leu as conversas deste período partindo de todas as categorias que já estão no Catálogo (inclusive "
        "as excluídas e as ainda sem decisão) e encontrou os temas abaixo, que não têm categoria. Eles NÃO foram "
        "usados nesta classificação — as conversas desses temas foram classificadas nas categorias existentes ou em "
        f'{CATEGORIA_SEM_CATALOGO}. Para usar uma sugestão, inclua a categoria no Catálogo com Decisão "Nova": '
        "ela passa a valer no próximo relatório."
    )
    por_id = dados.df.set_index("id_sessao_48h")
    for sugestao in dados.sugestoes:
        h = doc.add_heading(sugestao.nome, level=2)
        h.runs[0].font.color.rgb = _COR_ACCENT
        doc.add_paragraph(sugestao.descricao or "(sem descrição)")
        ids = [i for i in sugestao.ids_exemplo if i in por_id.index]
        if not ids:
            doc.add_paragraph().add_run("Sem exemplo disponível.").italic = True
            continue
        doc.add_paragraph().add_run("Exemplos:").bold = True
        tabela = doc.add_table(rows=1, cols=2)
        tabela.style = "Light Grid Accent 2"
        for cel, texto in zip(tabela.rows[0].cells, ["Conversa", "Classificada nesta rodada como"], strict=True):
            cel.paragraphs[0].add_run(texto).bold = True
        for id_sessao in ids:
            linha = por_id.loc[id_sessao]
            celulas = tabela.add_row().cells
            celulas[0].text = _trunca(linha["conversa_completa"])
            celulas[1].text = SEP_CELULA.join(linha["categorias"]) or "(falha na classificação)"
        doc.add_paragraph()


# ============================================================================
# Publicação
# ============================================================================


def publica_relatorio(
    drive: Resource, sheets: Resource, pasta_campanha_id: str, dados: DadosRelatorio, n_exemplos_por_categoria: int
) -> str:
    """Cria a subpasta da rodada e publica os arquivos. Devolve o link da subpasta."""
    categorias = monta_categorias(dados, n_exemplos_por_categoria)
    subpasta_id = garante_pasta(
        drive, pasta_campanha_id, nome_subpasta(dados.gerado_em, dados.data_inicio, dados.data_fim, dados.rapido)
    )

    buffer = io.BytesIO()
    gera_docx(dados, categorias).save(buffer)
    upload_bytes(drive, subpasta_id, ARQUIVO_RELATORIO, buffer.getvalue(), DOCX_MIME, converter_para=DRIVE_DOC_MIME)

    conversas = monta_conversas_df(dados.df, com_revisao=not dados.rapido)
    arquivo = cria_sheet(drive, sheets, subpasta_id, ARQUIVO_CONVERSAS, conversas)
    if not dados.rapido:
        aplica_lista_suspensa(sheets, arquivo["id"], COL_REVISAO, list(OPCOES_REVISAO), ate_linha=len(conversas) + 1)
    else:
        cria_sheet(
            drive, sheets, subpasta_id, ARQUIVO_CATALOGO_AUTOMATICO, monta_catalogo_automatico_df(dados, categorias)
        )
    return link_pasta(subpasta_id)
