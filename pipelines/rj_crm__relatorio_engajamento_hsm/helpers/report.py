"""The report: categories, margin of error, Word document and CSV files."""

from collections.abc import Sequence
from csv import writer
from datetime import date, datetime
from io import BytesIO, StringIO
from math import nan, sqrt

from docx import Document
from docx.document import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from pipelines.rj_crm__relatorio_engajamento_hsm.constants import (
    CATEGORY_COLUMNS,
    CHATBOT_DATASET,
    CHATBOT_TABLE,
    CORRECT,
    CSV_SEP,
    DETAIL_COLUMNS,
    MAX_EXAMPLE_CHARS,
    NOTE_FILES,
    NOTE_LLM,
    NOTE_PONTUAL,
    PROJECT_ID,
    QUERY_NOTE,
    Z_95,
)
from pipelines.rj_crm__relatorio_engajamento_hsm.helpers.text import slugify_hsm
from pipelines.rj_crm__relatorio_engajamento_hsm.types import Category, Classification, Example, Report

ACCENT = RGBColor(0xB5, 0x50, 0x2E)
CODE_FILL = "FBE9E2"
CODE_BORDER = "B5502E"
NOT_JUDGED = "não avaliada"


def wilson_margin(proportion: float, judged: int, size: int | None = None, z: float = Z_95) -> float:
    """Return the half width of the 95% Wilson interval of a proportion.

    `judged` is the number of observations.

    With `size`, the size of the category, the interval has the correction of a finite
    population. Without it a small category judged in full, a census with no sampling
    error, would show a margin over zero, which is wrong. Return NaN for no observation.
    """
    if judged <= 0:
        return nan

    if size is not None and judged >= size:
        return 0.0

    margin = z * sqrt(proportion * (1 - proportion) / judged + z**2 / (4 * judged**2)) / (1 + z**2 / judged)

    return margin * sqrt((size - judged) / (size - 1)) if size is not None and size > 1 else margin


def precision_of(category: Category) -> float | None:
    """Return the share of judged conversations that the judge found right.

    It gives nothing if none was judged.
    """
    return category.correct / category.judged if category.judged else None


def margin_of(category: Category) -> float | None:
    """Return the margin of error of the precision of a category, or nothing."""
    precision = precision_of(category)

    return None if precision is None else wilson_margin(precision, category.judged, category.count)


def describe_precision(category: Category) -> str:
    """Return the precision of a category as `87% ± 5% (n=20)`.

    A category that was not judged says so.
    """
    precision, margin = precision_of(category), margin_of(category)

    return (
        NOT_JUDGED if precision is None or margin is None else f"{precision:.0%} ± {margin:.0%} (n={category.judged})"
    )


def build_categories(
    items: Sequence[Classification], catalog: dict[str, str], examples_per_category: int
) -> list[Category]:
    """Return a category for each item name, the biggest first.

    Each has its numbers and examples.
    """
    by_name: dict[str, list[Classification]] = {}

    for item in items:
        by_name.setdefault(item.category, []).append(item)

    categories = [
        Category(
            name=name,
            description=catalog.get(name, ""),
            count=len(group),
            pct=100 * len(group) / len(items),
            judged=sum(1 for item in group if item.verdict.strip()),
            correct=sum(1 for item in group if item.verdict == CORRECT),
            examples=[
                Example(item.conversation.text, item.summary)
                for item in [item for item in group if item.conversation.text][:examples_per_category]
            ],
        )
        for name, group in sorted(by_name.items())
    ]

    return sorted(categories, key=lambda category: category.count, reverse=True)


def cut_text(text: str, limit: int = MAX_EXAMPLE_CHARS) -> str:
    """Return the text cut to the limit, with a mark when it was cut."""
    return text if len(text) <= limit else text[:limit].rstrip() + " […]"


def shade(paragraph: Paragraph, fill: str, border: str) -> None:
    """Give a paragraph a background and a border, as a block of code has."""
    properties = paragraph._p.get_or_add_pPr()  # pyright: ignore[reportPrivateUsage]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)
    borders = OxmlElement("w:pBdr")

    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "4")
        element.set(qn("w:color"), border)
        borders.append(element)

    properties.append(borders)


def add_code_block(document: WordDocument, text: str) -> None:
    """Add a block of code: monospace, with a line break for each line of the text."""
    paragraph = document.add_paragraph()
    lines = text.split("\n")

    for index, line in enumerate(lines):
        run = paragraph.add_run(line if line.strip() else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)

        if index < len(lines) - 1:
            run.add_break()

    shade(paragraph, CODE_FILL, CODE_BORDER)


def first_filled(values: Sequence[str | None]) -> str | None:
    """Return the first value that is not empty, or nothing."""
    return next((value for value in values if value), None)


def add_summary(document: WordDocument, report: Report, categories: Sequence[Category], hsm_text: str) -> None:
    """Add the summary: dispatches, engagement, conversations, judged and precision."""
    classified = len(report.items)
    judged = sum(category.judged for category in categories)
    correct = sum(category.correct for category in categories)
    document.add_heading("Resumo", level=1)
    document.add_paragraph(f"Total de disparos: {report.total_dispatches}")
    rate = 100 * report.engaged / report.total_dispatches if report.total_dispatches else 0.0
    document.add_paragraph(f"Sessões com resposta do cidadão: {report.engaged} ({rate:.1f}% dos disparos)")
    share = 100 * classified / report.engaged if report.engaged else 0.0
    document.add_paragraph(f"Conversas classificadas: {classified} ({share:.1f}% das respostas)")
    document.add_paragraph(f"Categorias identificadas: {len(categories)}")

    if classified:
        document.add_paragraph(
            f"Conversas avaliadas pelo juiz: {judged} ({100 * judged / classified:.1f}% das classificadas)"
        )

    line = document.add_paragraph()
    line.add_run("Precisão média da classificação (todas as categorias): ")

    if judged:
        precision = correct / judged
        line.add_run(f"{precision:.0%} ± {wilson_margin(precision, judged, classified):.0%} (n={judged})").bold = True
    else:
        line.add_run("não avaliada")

    if hsm_text:
        document.add_paragraph().add_run("HSM enviado:").bold = True
        document.add_paragraph(style="Intense Quote").add_run(cut_text(hsm_text))


def add_notes(document: WordDocument, csv_names: Sequence[str]) -> None:
    """Add the notes on how to read the report."""
    names = " e ".join(f"'{name}'" for name in csv_names)
    document.add_heading("Observações importantes", level=1)
    document.add_paragraph(NOTE_PONTUAL)
    document.add_paragraph(NOTE_LLM)
    document.add_paragraph(NOTE_FILES.format(names=names, separator=CSV_SEP))


def add_categories(document: WordDocument, categories: Sequence[Category]) -> None:
    """Add the table of the categories and, for each, a description and examples."""
    document.add_heading("Categorias", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 2"

    for cell, text in zip(
        table.rows[0].cells, ["Categoria", "Qtd.", "% do engajamento", "Precisão (LLM juiz)"], strict=True
    ):
        cell.paragraphs[0].add_run(text).bold = True

    for category in categories:
        cells = table.add_row().cells
        cells[0].text = category.name
        cells[1].text = str(category.count)
        cells[2].text = f"{category.pct:.1f}%"
        cells[3].text = describe_precision(category)

    document.add_paragraph()

    for category in categories:
        heading = document.add_heading(category.name, level=2)
        heading.runs[0].font.color.rgb = ACCENT
        document.add_paragraph(category.description or "(sem descrição registrada no catálogo)")
        line = document.add_paragraph()
        line.add_run(f"{category.count} conversa(s) · {category.pct:.1f}% do engajamento · precisão: ")
        precision, margin = precision_of(category), margin_of(category)

        if precision is None or margin is None:
            line.add_run(NOT_JUDGED)
        else:
            line.add_run(f"{precision:.0%} ± {margin:.0%} (amostra de {category.judged})").bold = True

        add_examples(document, category)


def add_examples(document: WordDocument, category: Category) -> None:
    """Add the examples of a category as a table of conversation and summary."""
    if not category.examples:
        document.add_paragraph().add_run("Sem exemplo disponível.").italic = True
        return

    document.add_paragraph().add_run("Exemplos:").bold = True
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 2"

    for cell, text in zip(table.rows[0].cells, ["Conversa", "Resumo"], strict=True):
        cell.paragraphs[0].add_run(text).bold = True

    for example in category.examples:
        cells = table.add_row().cells
        cells[0].text = cut_text(example.conversation)
        cells[1].text = cut_text(example.summary) if example.summary else "(sem resumo)"

    document.add_paragraph()


def build_docx(
    report: Report, categories: Sequence[Category], sql: str, csv_names: Sequence[str], now: datetime
) -> bytes:
    """Return the Word document of a report.

    It has the summary, notes, the query, the categories and examples.

    `sql` is the query that read the conversations, with its values in place, for anyone
    to reproduce.
    """
    document = Document()
    first = report.items[0].conversation if report.items else None
    items = [item.conversation for item in report.items]
    document.add_heading(f"Relatório de Engajamento — {report.dispatch.hsm_name}", level=0)
    meta = document.add_paragraph()
    campaign, axis = first_filled([item.campaign for item in items]), first_filled([item.axis for item in items])
    meta.add_run(f"Campanha: {campaign or '—'}    ·    Eixo: {axis or '—'}").italic = True
    document.add_paragraph(f"Gerado em {now:%d/%m/%Y %H:%M}")
    add_summary(document, report, categories, first.hsm_text if first else "")
    add_notes(document, csv_names)
    document.add_heading("Query utilizada na análise", level=1)
    document.add_paragraph(QUERY_NOTE.format(table=f"{PROJECT_ID}.{CHATBOT_DATASET}.{CHATBOT_TABLE}"))
    add_code_block(document, sql)
    add_categories(document, categories)
    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def to_csv(columns: Sequence[str], rows: Sequence[Sequence[str | int | float | None]]) -> bytes:
    """Return rows as CSV with `;` as the separator, and nothing as an empty cell."""
    buffer = StringIO()
    out = writer(buffer, delimiter=CSV_SEP, lineterminator="\n")
    out.writerow(columns)
    out.writerows(["" if cell is None else cell for cell in row] for row in rows)

    return buffer.getvalue().encode("utf-8")


def categories_csv(categories: Sequence[Category]) -> bytes:
    """Return the CSV of the categories, one line each.

    It has the share, judged count, precision and margin.
    """
    rows = [
        (
            category.name,
            category.description,
            category.count,
            round(category.pct, 1),
            category.judged,
            None if (precision := precision_of(category)) is None else round(precision, 4),
            None if (margin := margin_of(category)) is None else round(margin, 4),
        )
        for category in categories
    ]

    return to_csv(CATEGORY_COLUMNS, rows)


def detail_csv(items: Sequence[Classification]) -> bytes:
    """Return the CSV of the conversations: one line with its category and verdict."""
    rows = [
        (
            item.conversation.session_id,
            item.conversation.cpf,
            item.conversation.phone,
            item.category,
            item.justification,
            item.summary,
            item.verdict,
            item.expected_category,
            item.judge_note,
        )
        for item in items
    ]

    return to_csv(DETAIL_COLUMNS, rows)


def report_slug(hsm_name: str, reference: date) -> str:
    """Return the prefix of the report files of an HSM for a day, which has no time.

    It names the files, and it tells whether the report exists, so a run repeats with no
    duplicate.
    """
    return f"relatorio_engajamento_{slugify_hsm(hsm_name)}__geracao_{reference.isoformat()}"
